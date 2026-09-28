#!/usr/bin/env python3
"""Expoline QA Test 24: mandatory service charge — Bali Hai 18% treatment.

Covers phase 4:
  - charge applies at/above the guest threshold, never below
  - CA tax base: subtotal + surcharge + mandatory service charge (CDTFA
    Publication 22, Jan 2025; Annotation 550.0740) — mandatory charges are
    part of taxable gross receipts; voluntary tips are never taxed
  - service charge is NEVER mixed with tip lines: separate totals fields,
    separate report columns, Tips report excludes it entirely
  - percentage + threshold configurable per site (manager session + live
    manager PIN + audit log with before/after); server enforces on every
    recompute; invalid values rejected
  - reports (sales/tax/tips/payouts/shift/labor) show it separately and correctly;
    payouts carry an explicit informational service_charge_cents (inside card
    volume, never double-counted); labor states it is house revenue, not
    wages/tips; .xlsx/.csv/.pdf/.docx exports carry the notes incl. the
    "confirm with your accountant" disclaimer
  - KDS intentionally carries no financial lines (asserted, not just noted)
  - printable guest receipt exists in the frontend (printReceipt)
  - split stays blocked on charged checks, with a config-driven message

Date-agnostic: uses the site date from /api/config.
"""
import io, json, os, sys, urllib.request, urllib.error, zipfile

BASE = os.environ.get("EXPOLINE_BASE", os.environ.get("EXPLOINE_BASE", "http://localhost:4328"))
S, M = "1111", "2580"

def api(method, path, token=None, body=None, raw=False):
    req = urllib.request.Request(BASE + path, method=method,
        data=json.dumps(body).encode() if body is not None else None,
        headers={"Content-Type": "application/json"})
    if token: req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req) as resp:
            return (resp.status, resp.read()) if raw else json.loads(resp.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        return (e.code, e.read()) if raw else (_ for _ in ()).throw(e)

def login(pin):
    return api("POST", "/api/auth/login", None, {"pin": pin})["token"]

ST = login(S); MT = login(M)
checks, fails = 0, []
def ok(cond, name, detail=""):
    global checks; checks += 1
    print(f"  {'\u2713' if cond else '\u2717'} {name}" + (f" {detail}" if detail and not cond else ""))
    if not cond: fails.append(f"FAIL: {name} {detail}")

SITE_DATE = api("GET", "/api/config", MT)["site_date"]
ITEM = 71          # Lava Slide, 1400¢ (seed)
TBL = 10   # bumped per check — one open check per table
def new_check(guests, tab):
    global TBL; TBL += 1
    return api("POST", "/api/checks", ST, {"table_id": TBL, "guest_count": guests, "tab_name": tab})
def add2(c):
    return api("POST", f"/api/checks/{c['id']}/items", ST, {"menu_item_id": ITEM, "seat": 1, "qty": 2, "modifiers": []})
def getc(c):
    return api("GET", f"/api/checks/{c['id']}", ST)

print("== 24a: threshold boundary (default 18% on 8+) ==")
c7 = new_check(7, "QA svc 7-top"); add2(c7)
t7 = getc(c7)
ok(t7["service_charge_cents"] == 0, "7 guests -> no service charge", str(t7["service_charge_cents"]))
exp_tax7 = round((2800 + 140) * 0.0775)
ok(t7["tax_cents"] == exp_tax7, "7-top tax on sub+sur only", f"{t7['tax_cents']} vs {exp_tax7}")
ok(t7["total_cents"] == 2800 + 140 + exp_tax7, "7-top total", str(t7["total_cents"]))

c8 = new_check(8, "QA svc 8-top"); add2(c8)
t8 = getc(c8)
ok(t8["service_charge_cents"] == 504, "8 guests -> 18% charge = 504", str(t8["service_charge_cents"]))
exp_tax8 = round((2800 + 140 + 504) * 0.0775)   # CA: mandatory charge IS in the taxable base
ok(exp_tax8 == 267, "hand-computed tax 267", str(exp_tax8))
ok(t8["tax_cents"] == 267, "8-top tax INCLUDES service charge (CDTFA)", f"{t8['tax_cents']} vs 267")
ok(t8["total_cents"] == 3711, "8-top total 2800+140+504+267", str(t8["total_cents"]))
ok(t8["totals"]["service_charge"] == 504 and t8["totals"]["tax"] == 267, "totals block mirrors row")

print("== 24b: split blocked, message is config-driven ==")
st, body = api("POST", f"/api/checks/{c8['id']}/split", ST, {"mode": "even", "parts": 2}, raw=True)
err = json.loads(body.decode())["error"]
ok(st == 400, "split on charged check -> 400", str(st))
ok("18%" in err, "split message names the 18% charge", err)

print("== 24c: config endpoints ==")
cfg = api("GET", "/api/admin/service-charge/config", MT)
ok(cfg["current"]["service_charge_pct"] == 0.18, "current pct 0.18", str(cfg["current"]))
ok(cfg["current"]["service_charge_min_guests"] == 8, "current threshold 8", str(cfg["current"]))
ok(cfg["defaults"]["service_charge_pct"] == 0.18, "defaults reported", str(cfg["defaults"]))
ok("accountant" in cfg.get("note", ""), "config GET carries accountant note", cfg.get("note", "")[:40])

st, _ = api("PUT", "/api/admin/service-charge/config", ST,
            {"key": "service_charge_pct", "value": 0.2, "manager_pin": M}, raw=True)
ok(st == 403, "server role cannot change config -> 403", str(st))
st, _ = api("PUT", "/api/admin/service-charge/config", MT,
            {"key": "service_charge_pct", "value": 0.2, "manager_pin": "0000"}, raw=True)
ok(st == 403, "wrong manager PIN -> 403", str(st))
st, _ = api("PUT", "/api/admin/service-charge/config", MT,
            {"key": "nope", "value": 1, "manager_pin": M}, raw=True)
ok(st == 400, "unknown key -> 400", str(st))
st, _ = api("PUT", "/api/admin/service-charge/config", MT,
            {"key": "service_charge_pct", "value": 0.9, "manager_pin": M}, raw=True)
ok(st == 400, "pct > 50% rejected -> 400", str(st))
st, _ = api("PUT", "/api/admin/service-charge/config", MT,
            {"key": "service_charge_min_guests", "value": -1, "manager_pin": M}, raw=True)
ok(st == 400, "negative threshold rejected -> 400", str(st))
st, _ = api("PUT", "/api/admin/service-charge/config", MT,
            {"key": "service_charge_min_guests", "value": 2.5, "manager_pin": M}, raw=True)
ok(st == 400, "non-integer threshold rejected -> 400", str(st))

print("== 24d: pct change takes effect server-side ==")
r = api("PUT", "/api/admin/service-charge/config", MT,
        {"key": "service_charge_pct", "value": 0.2, "manager_pin": M})
ok(r["value"] == 0.2 and r["current"]["service_charge_pct"] == 0.2, "pct set to 20%", str(r))
c20 = new_check(8, "QA svc 20pct"); add2(c20)
t20 = getc(c20)
ok(t20["service_charge_cents"] == 560, "20% of 2800 = 560", str(t20["service_charge_cents"]))
ok(t20["tax_cents"] == round((2800 + 140 + 560) * 0.0775), "tax follows new pct", str(t20["tax_cents"]))
st, body = api("POST", f"/api/checks/{c20['id']}/split", ST, {"mode": "even", "parts": 2}, raw=True)
ok(st == 400 and "20%" in json.loads(body.decode())["error"], "split message tracks new pct", body.decode()[:80])
aud = api("GET", "/api/admin/service-charge/audit?limit=5", MT)
latest = aud[0]
ok(latest["action"] == "config_change", "audit row recorded", str(latest["action"]))
before, after = json.loads(latest["before_json"]), json.loads(latest["after_json"])
ok(before["service_charge_pct"] == 0.18 and after["service_charge_pct"] == 0.2,
   "audit before/after pct 0.18 -> 0.2", f"{before['service_charge_pct']} -> {after['service_charge_pct']}")
ok(after.get("approver"), "audit records PIN approver", str(after.get("approver")))
# restore
api("PUT", "/api/admin/service-charge/config", MT, {"key": "service_charge_pct", "value": 0.18, "manager_pin": M})
ok(api("GET", "/api/admin/service-charge/config", MT)["current"]["service_charge_pct"] == 0.18, "pct restored to 18%")

print("== 24e: threshold change + disable ==")
api("PUT", "/api/admin/service-charge/config", MT, {"key": "service_charge_min_guests", "value": 4, "manager_pin": M})
ok(getc(c7)["service_charge_cents"] == 504, "threshold 4 -> 7-top now charged", str(getc(c7)["service_charge_cents"]))
api("PUT", "/api/admin/service-charge/config", MT, {"key": "service_charge_min_guests", "value": 0, "manager_pin": M})
ok(api("GET", "/api/admin/service-charge/config", MT)["disabled"] is True, "threshold 0 -> disabled flag")
ok(getc(c8)["service_charge_cents"] == 0, "disabled -> 8-top no charge")
api("PUT", "/api/admin/service-charge/config", MT, {"key": "service_charge_min_guests", "value": 8, "manager_pin": M})
ok(getc(c8)["service_charge_cents"] == 504, "threshold restored -> 8-top charged again")
aud2 = api("GET", "/api/admin/service-charge/audit?limit=20", MT)
ok(len(aud2) >= 5, "audit trail holds all changes", str(len(aud2)))

print("== 24f: service charge never mixed with tips ==")
# fresh 8-top at default config; pay with a tip
cp = new_check(8, "QA svc tip-sep"); add2(cp); getc(cp)
p = api("POST", f"/api/checks/{cp['id']}/payments", ST,
        {"method": "card_demo", "amount_cents": 3711, "tip_cents": 1000, "brand": "Visa", "last4": "4242"})
ok(p["payment"]["tip_cents"] == 1000, "tip stored on payment", str(p["payment"]["tip_cents"]))
ok(p["check"]["totals"]["service_charge"] == 504, "check totals keep service charge", str(p["check"]["totals"]["service_charge"]))
ok("service" not in json.dumps(p["payment"]).lower(), "payment row has no service-charge field")
api("POST", f"/api/checks/{cp['id']}/close", ST)

rep = lambda r, **kw: api("GET", f"/api/finance/reports/{r}?format=json&period=day&date={SITE_DATE}", MT, **kw)
tips = rep("tips")
tt = tips["totals"]["total_tips_cents"]
ok(tt == 1000, "tips report totals = only the 1000 tip", str(tt))
ok(not any("service" in json.dumps(row).lower() for row in tips["rows"]), "tips rows carry no service data")
sales = rep("sales")
st8 = sales["totals"]
ok(st8["service_cents"] == 504, "sales report: service charge separate column = 504", str(st8["service_cents"]))
ok(st8["tips_cents"] == 1000, "sales report: tips separate = 1000", str(st8["tips_cents"]))
ok(st8["net_cents"] == st8["gross_cents"] + st8["surcharge_cents"] + st8["service_cents"] - st8["comp_cents"],
   "net = gross + surcharge + service - comps", str(st8["net_cents"]))
tax = rep("tax")
tx = tax["totals"]
ok(tx["taxable_cents"] == 2800 + 140 + 504, "tax report taxable includes service charge", str(tx["taxable_cents"]))
ok(tx["tax_cents"] == 267, "tax report tax = 267", str(tx["tax_cents"]))
payouts = rep("payouts")
po = payouts["totals"]
ok(po["expected_payout_cents"] == po["card_volume_cents"] - po["refunds_cents"] - po["stripe_fees_cents"],
   "payout math untouched by service charge", str(po))
po_day = api("GET", f"/api/finance/payouts?date={SITE_DATE}", MT)
ok(po_day["tips_cents"] == 1000, "payouts: tips collected separately", str(po_day["tips_cents"]))
shift = api("GET", f"/api/finance/shift?date={SITE_DATE}", MT)
ok(shift["service_charge_cents"] == 504, "shift report: service_charge_cents = 504", str(shift.get("service_charge_cents")))
ok(shift["tips_cents"] == 1000, "shift report: tips separate = 1000", str(shift["tips_cents"]))

print("== 24g: exports carry service charge separately + notes ==")
def exp_bytes(r, fmt):
    st, b = api("GET", f"/api/finance/reports/{r}?format={fmt}&period=day&date={SITE_DATE}", MT, raw=True)
    ok(st == 200, f"{r}.{fmt} -> 200", str(st))
    return b
x = exp_bytes("sales", "xlsx")
ok(x[:2] == b"PK", "xlsx magic bytes", str(x[:2]))
znames = zipfile.ZipFile(io.BytesIO(x)).namelist()
ok("xl/worksheets/sheet2.xml" in znames or len(znames) > 3, "xlsx has multiple sheets (Report + Notes)", str(len(znames)))
wb = zipfile.ZipFile(io.BytesIO(x)).read("xl/workbook.xml").decode()
ok("Notes" in wb, "xlsx includes Notes sheet", wb[:200])
pdf = exp_bytes("sales", "pdf")
ok(pdf[:4] == b"%PDF", "pdf magic bytes", str(pdf[:4]))
# pdf-lib Flate-encodes streams and hex-encodes text: decompress + decode hex strings
import re as _re, zlib as _zlib, binascii as _ba
_ptext = b""
for _m in _re.finditer(rb"stream\r?\n(.*?)endstream", pdf, _re.S):
    try: _ptext += _zlib.decompress(_m.group(1))
    except Exception: pass
_pdfwords = b" ".join(_ba.unhexlify(h) for h in _re.findall(rb"<([0-9a-fA-F]+)>", _ptext) if len(h) % 2 == 0)
ok(b"Service charge" in _pdfwords, "pdf contains Service charge column")
ok(b"CDTFA" in _pdfwords and b"accountant" in _pdfwords, "pdf notes cite CDTFA + accountant disclaimer")
docx = exp_bytes("sales", "docx")
ok(docx[:2] == b"PK", "docx magic bytes", str(docx[:2]))
ddoc = zipfile.ZipFile(io.BytesIO(docx)).read("word/document.xml").decode()
ok("Service charge" in ddoc, "docx contains Service charge column")
ok("CDTFA" in ddoc and "accountant" in ddoc, "docx notes cite CDTFA + accountant disclaimer")
csv = exp_bytes("sales", "csv").decode()
ok("Service charge" in csv, "csv contains Service charge column")
ok("CDTFA" in csv and "accountant" in csv, "csv carries CDTFA note + accountant disclaimer")
tcsv = exp_bytes("tips", "csv").decode()
ok("NOT a tip" in tcsv, "tips csv notes state service charge is NOT a tip")
ok("Service charge" not in tcsv.splitlines()[0], "tips csv has no service-charge column")
pdocx = exp_bytes("payouts", "docx")
pddoc = zipfile.ZipFile(io.BytesIO(pdocx)).read("word/document.xml").decode()
ok("Service charge is restaurant revenue" in pddoc, "payout export notes name service-charge treatment")
tdocx = exp_bytes("tax", "docx")
tddoc = zipfile.ZipFile(io.BytesIO(tdocx)).read("word/document.xml").decode()
ok("550.0740" in tddoc, "tax export notes cite Annotation 550.0740")

print("== 24h: split message after restore ==")
st, body = api("POST", f"/api/checks/{c8['id']}/split", ST, {"mode": "even", "parts": 2}, raw=True)
ok(st == 400 and "18%" in json.loads(body.decode())["error"], "split on charged check -> 400 naming 18%", body.decode()[:80])

print("== 24i: payout reconciliation carries service charge separately ==")
po_day2 = api("GET", f"/api/finance/payouts?date={SITE_DATE}", MT)
ok(po_day2["service_charge_cents"] == 504, "payout endpoint: service_charge_cents = 504", str(po_day2.get("service_charge_cents")))
ok(po_day2["expected_payout_cents"] == po_day2["card_volume_cents"] - po_day2["refunds_cents"] - po_day2["stripe_fees_cents"],
   "payout math unchanged: service charge already inside card volume, not double-counted")
ok("service charge" in (po_day2.get("service_charge_note") or "").lower() and "tip" in (po_day2.get("service_charge_note") or "").lower(),
   "payout endpoint carries service-charge informational note")
px = exp_bytes("payouts", "xlsx")
import openpyxl as _oxl
pwb = _oxl.load_workbook(io.BytesIO(px), read_only=True, data_only=True)
prep = pwb["Report"]
headers = [c.value for c in prep[4]]
ok("Service charge (info)" in headers, "payouts xlsx has separate Service charge (info) column", str(headers))
totrow = [c.value for c in prep[prep.max_row]]
ok(totrow[0] == "Total" and totrow[headers.index("Service charge (info)")] == 504,
   "payouts xlsx totals row: service charge = 504", str(totrow))
pnotes = " ".join(str(c.value) for row in pwb["Notes"].iter_rows() for c in row if c.value)
ok("informational" in pnotes.lower() and "card volume" in pnotes.lower(),
   "payouts xlsx notes mark the column informational (inside card volume)")
pcsv = exp_bytes("payouts", "csv").decode("utf-8")
ok("Service charge (info)" in pcsv.splitlines()[0], "payouts csv header has Service charge (info) column")
ok("informational" in pcsv.lower() and "card volume" in pcsv.lower(), "payouts csv carries the notes incl. disclaimer")

print("== 24j: labor report states house-revenue treatment ==")
lx = exp_bytes("labor", "xlsx")
lwb = _oxl.load_workbook(io.BytesIO(lx), read_only=True, data_only=True)
lrep = lwb["Report"]
lheaders = [c.value for c in lrep[4]]
ok(not any("service" in str(h).lower() for h in lheaders if h), "labor xlsx has no service-charge column", str(lheaders))
lnotes = " ".join(str(c.value) for row in lwb["Notes"].iter_rows() for c in row if c.value)
ok("house revenue" in lnotes.lower(), "labor notes: service charges are house revenue")
ok("not wages" in lnotes.lower(), "labor notes: not wages")
lcsv = exp_bytes("labor", "csv").decode("utf-8")
ok("house revenue" in lcsv.lower(), "labor csv carries the house-revenue note")

print("== 24k: KDS carries no financial lines by design ==")
ck = new_check(8, "QA svc kds"); add2(ck)
api("POST", f"/api/checks/{ck['id']}/send", ST)
KT = login("2222")
tk = api("GET", "/api/kds/tickets", KT)
ok(isinstance(tk, list) and len(tk) > 0, "KDS ticket exists for sent items")
tj = json.dumps(tk).lower()
for _key in ["subtotal", "surcharge", "service_charge", "tax_cents", "tip_cents", "total_cents"]:
    ok(_key not in tj, f"KDS ticket JSON carries no '{_key}'")

print("== 24l: printable receipt exists in the frontend ==")
import os as _os
_appjs = open(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "public", "app.js")).read()
ok("function printReceipt" in _appjs, "frontend has printReceipt()")
ok("Print receipt" in _appjs, "pay view has a Print receipt button")
ok("NOT a tip" in _appjs, "receipt labels the service charge as NOT a tip")
ok("accountant" in _appjs, "receipt carries the accountant disclaimer")

print(f"\nTest 24: {checks} assertions, {len(fails)} failures")
for f in fails: print(f)
sys.exit(1 if fails else 0)
